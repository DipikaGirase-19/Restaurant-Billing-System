import tkinter as tk
from tkinter import messagebox, ttk
from openpyxl import Workbook, load_workbook
import os
from datetime import datetime

# --- CONFIG ---
EXCEL_FILE = "Restaurant_Bill_Data.xlsx"
HEADERS = ["Bill ID", "Customer Name", "Phone", "Food Item", "Quantity", "Price", "Total", "Date"]
MENU = {"Pizza": 150, "Burger": 100, "Sandwich": 80, "French Fries": 70, "Cold Drink": 50}

def init_excel():
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "Restaurant Bills"
        ws.append(HEADERS)
        wb.save(EXCEL_FILE)

def get_records():
    if not os.path.exists(EXCEL_FILE): return []
    wb = load_workbook(EXCEL_FILE)
    ws = wb.active
    return list(ws.iter_rows(min_row=2, values_only=True))

init_excel()

# 1. LOGIN PAGE - Requirement 1
def login_page():
    root = tk.Tk()
    root.title("Login Page - Restaurant System")
    root.geometry("400x320")
    root.config(bg="#eef2f7")

    tk.Label(root, text="RESTAURANT BILLING SYSTEM", font=("Arial", 16, "bold"), bg="#eef2f7").pack(pady=20)
    tk.Label(root, text="Login Page", font=("Arial", 11), bg="#eef2f7").pack()

    tk.Label(root, text="Username", bg="#eef2f7").pack(pady=(20,2))
    u_entry = tk.Entry(root, width=28, font=("Arial", 11))
    u_entry.pack()
    tk.Label(root, text="Password", bg="#eef2f7").pack(pady=(10,2))
    p_entry = tk.Entry(root, width=28, show="*", font=("Arial", 11))
    p_entry.pack()

    def check_login():
        if u_entry.get() == "admin" and p_entry.get() == "1234":
            root.destroy()
            dashboard()
        else:
            messagebox.showerror("Error", "Invalid Username or Password!")

    tk.Button(root, text="Login", bg="#0b5ed7", fg="white", width=15, font=("Arial", 11, "bold"), command=check_login).pack(pady=25)
    tk.Label(root, text="Hint: admin / 1234", fg="gray", bg="#eef2f7").pack()
    root.mainloop()

# 2. DASHBOARD - Requirement 2
def dashboard():
    dash = tk.Tk()
    dash.title("Dashboard - Restaurant Billing System")
    dash.geometry("850x500")
    dash.config(bg="white")

    top = tk.Frame(dash, bg="#0b2c4d", height=65)
    top.pack(fill="x")
    tk.Label(top, text="RESTAURANT BILLING SYSTEM", font=("Arial", 20, "bold"), bg="#0b2c4d", fg="white").pack(side="left", padx=20, pady=15)
    # Logout option - Requirement 1.3
    tk.Button(top, text="Logout", bg="#dc3545", fg="white", font=("Arial", 10, "bold"), command=lambda:[dash.destroy(), login_page()]).pack(side="right", padx=20)

    tk.Label(dash, text="Dashboard - Select Option", font=("Arial", 16, "bold"), bg="white").pack(pady=20)
    f = tk.Frame(dash, bg="white")
    f.pack()

    tk.Button(f, text="1. Add Record", width=22, height=2, bg="#198754", fg="white", font=("Arial", 13, "bold"), command=lambda:[dash.destroy(), add_record()]).grid(row=0, column=0, padx=15, pady=10)
    tk.Button(f, text="2. View Records", width=22, height=2, bg="#0d6efd", fg="white", font=("Arial", 13, "bold"), command=lambda:[dash.destroy(), view_records()]).grid(row=0, column=1, padx=15, pady=10)
    tk.Button(f, text="3. Update / Delete", width=22, height=2, bg="#ffc107", font=("Arial", 13, "bold"), command=lambda:[dash.destroy(), view_records()]).grid(row=1, column=0, padx=15, pady=10)
    tk.Button(f, text="4. Exit", width=22, height=2, bg="#6c757d", fg="white", font=("Arial", 13, "bold"), command=dash.destroy).grid(row=1, column=1, padx=15, pady=10)

    dash.mainloop()

# 3. ADD RECORD - Requirement 3
def add_record():
    win = tk.Tk()
    win.title("Add Record")
    win.geometry("750x500")
    win.config(bg="white")
    tk.Label(win, text="Add New Bill - Tkinter Form", font=("Arial", 16, "bold"), bg="white").pack(pady=15)

    form = tk.Frame(win, bg="white")
    form.pack()

    tk.Label(form, text="Customer Name", bg="white").grid(row=0, column=0, sticky="w", pady=8)
    name_e = tk.Entry(form, width=25); name_e.grid(row=0, column=1, pady=8)
    tk.Label(form, text="Phone", bg="white").grid(row=1, column=0, sticky="w", pady=8)
    phone_e = tk.Entry(form, width=25); phone_e.grid(row=1, column=1, pady=8)
    tk.Label(form, text="Food Item", bg="white").grid(row=2, column=0, sticky="w", pady=8)
    item_cb = ttk.Combobox(form, values=list(MENU.keys()), width=23); item_cb.grid(row=2, column=1, pady=8); item_cb.set("Pizza")
    tk.Label(form, text="Quantity", bg="white").grid(row=3, column=0, sticky="w", pady=8)
    qty_e = tk.Entry(form, width=25); qty_e.grid(row=3, column=1, pady=8)

    price_l = tk.Label(form, text=f"Price: ₹{MENU['Pizza']}", fg="blue", font=("Arial", 11, "bold"), bg="white")
    price_l.grid(row=2, column=2, padx=15)
    def change_price(e): price_l.config(text=f"Price: ₹{MENU[item_cb.get()]}")
    item_cb.bind("<<ComboboxSelected>>", change_price)

    # Clear/Reset - Requirement 7
    def clear_fields():
        name_e.delete(0, tk.END); phone_e.delete(0, tk.END); qty_e.delete(0, tk.END); item_cb.set("Pizza"); price_l.config(text=f"Price: ₹{MENU['Pizza']}")

    def save_to_excel():
        if not name_e.get() or not qty_e.get():
            messagebox.showerror("Error", "Fill all fields!"); return
        try:
            qty = int(qty_e.get()); assert qty > 0
        except:
            messagebox.showerror("Error", "Enter valid Quantity!"); return

        price = MENU[item_cb.get()]; total = qty * price
        bid = int(datetime.now().timestamp())
        date = datetime.now().strftime("%d-%m-%Y %H:%M")

        wb = load_workbook(EXCEL_FILE); ws = wb.active
        ws.append([bid, name_e.get(), phone_e.get(), item_cb.get(), qty, price, total, date])
        wb.save(EXCEL_FILE)
        messagebox.showinfo("Success", f"Saved to Excel!\nTotal ₹{total}")
        clear_fields()

    bf = tk.Frame(win, bg="white"); bf.pack(pady=20)
    tk.Button(bf, text="Save to Excel", bg="#198754", fg="white", width=15, command=save_to_excel).grid(row=0, column=0, padx=10)
    tk.Button(bf, text="Clear/Reset", bg="gray", fg="white", width=15, command=clear_fields).grid(row=0, column=1, padx=10)
    tk.Button(bf, text="Back to Dashboard", bg="#0b2c4d", fg="white", width=18, command=lambda:[win.destroy(), dashboard()]).grid(row=0, column=2, padx=10)
    win.mainloop()

# 4,5,6 - VIEW, UPDATE, DELETE
def view_records():
    win = tk.Tk()
    win.title("View Records - Treeview")
    win.geometry("1050x600")
    tk.Label(win, text="View / Update / Delete Records", font=("Arial", 15, "bold")).pack(pady=10)

    sf = tk.Frame(win); sf.pack()
    tk.Label(sf, text="Search by Customer:").pack(side="left")
    s_entry = tk.Entry(sf, width=30); s_entry.pack(side="left", padx=10)

    tree = ttk.Treeview(win, columns=HEADERS, show="headings", height=14)
    for h in HEADERS:
        tree.heading(h, text=h); tree.column(h, width=115)
    tree.pack(padx=10, pady=10, fill="both", expand=True)

    def load(filter_text=""):
        tree.delete(*tree.get_children())
        for r in get_records():
            if filter_text.lower() in str(r[1]).lower() or filter_text == "":
                tree.insert("", tk.END, values=r)

    tk.Button(sf, text="Search", command=lambda: load(s_entry.get())).pack(side="left")
    load()

    edit = tk.LabelFrame(win, text="Update Selected Record", padx=10, pady=10)
    edit.pack(fill="x", padx=10)

    tk.Label(edit, text="Customer").grid(row=0, column=0); n_e = tk.Entry(edit, width=15); n_e.grid(row=0, column=1, padx=5)
    tk.Label(edit, text="Phone").grid(row=0, column=2); p_e = tk.Entry(edit, width=12); p_e.grid(row=0, column=3, padx=5)
    tk.Label(edit, text="Item").grid(row=0, column=4); i_cb = ttk.Combobox(edit, values=list(MENU.keys()), width=12); i_cb.grid(row=0, column=5, padx=5)
    tk.Label(edit, text="Qty").grid(row=0, column=6); q_e = tk.Entry(edit, width=8); q_e.grid(row=0, column=7, padx=5)

    def on_select(e):
        sel = tree.selection()
        if not sel: return
        v = tree.item(sel[0])['values']
        n_e.delete(0, tk.END); n_e.insert(0, v[1]); p_e.delete(0, tk.END); p_e.insert(0, v[2]); i_cb.set(v[3]); q_e.delete(0, tk.END); q_e.insert(0, v[4])
    tree.bind("<<TreeviewSelect>>", on_select)

    def update_rec():
        sel = tree.selection()
        if not sel: messagebox.showerror("Error", "Select a record!"); return
        bid = tree.item(sel[0])['values'][0]
        try:
            qty = int(q_e.get()); price = MENU[i_cb.get()]; total = qty*price
        except: messagebox.showerror("Error", "Invalid Qty!"); return
        wb = load_workbook(EXCEL_FILE); ws = wb.active
        for row in ws.iter_rows(min_row=2):
            if row[0].value == bid:
                row[1].value=n_e.get(); row[2].value=p_e.get(); row[3].value=i_cb.get(); row[4].value=qty; row[5].value=price; row[6].value=total; break
        wb.save(EXCEL_FILE); load(); messagebox.showinfo("Success", "Updated in Excel!")

    def delete_rec():
        sel = tree.selection()
        if not sel: messagebox.showerror("Error", "Select a record!"); return
        if not messagebox.askyesno("Confirmation", "Delete this record?"): return # Confirmation before deletion
        bid = tree.item(sel[0])['values'][0]
        wb = load_workbook(EXCEL_FILE); ws = wb.active
        for row in ws.iter_rows(min_row=2):
            if row[0].value == bid:
                ws.delete_rows(row[0].row); break
        wb.save(EXCEL_FILE); load(); messagebox.showinfo("Deleted", "Deleted from Excel!")

    tk.Button(edit, text="Update in Excel", bg="#ffc107", font=("Arial", 10, "bold"), command=update_rec).grid(row=0, column=8, padx=15)
    tk.Button(edit, text="Delete from Excel", bg="red", fg="white", font=("Arial", 10, "bold"), command=delete_rec).grid(row=0, column=9, padx=5)
    tk.Button(win, text="Back to Dashboard", command=lambda:[win.destroy(), dashboard()]).pack(pady=5)
    win.mainloop()

login_page()